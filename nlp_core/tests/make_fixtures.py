"""Regenerate the PDF / DOCX / scanned-PDF fixtures from sample_supply.txt.  Run: python tests/make_fixtures.py"""
from pathlib import Path

HERE = Path(__file__).resolve().parent / "fixtures"
text = (HERE / "sample_supply.txt").read_text(encoding="utf-8")

# --- DOCX with plain numbered text
import docx
d = docx.Document()
for line in text.split("\n"):
    if line.strip():
        d.add_paragraph(line)
d.save(HERE / "sample_supply.docx")

# --- DOCX with Word auto-numbering (numbers are NOT in the text)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
d2 = docx.Document()
for heading in ["Payment", "Termination"]:
    p = d2.add_paragraph(heading + " Buyer shall pay within 10 days." if heading == "Payment" else heading + " Either party may terminate.")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "1")
    numPr.append(il); numPr.append(ni); pPr.append(numPr)
d2.save(HERE / "autonumbered.docx")

# --- text PDF
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
c = canvas.Canvas(str(HERE / "sample_supply.pdf"), pagesize=A4)
y = 800
for line in text.split("\n"):
    if y < 60:
        c.showPage(); y = 800
    c.drawString(50, y, line[:110]); y -= 16
c.save()

# --- "scanned" PDF: pages with only a drawn rectangle, no text layer
c = canvas.Canvas(str(HERE / "scanned.pdf"), pagesize=A4)
for _ in range(2):
    c.rect(50, 50, 400, 600, fill=1); c.showPage()
c.save()
print("fixtures written to", HERE)
